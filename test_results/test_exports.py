# [WFGY] Zone: TEST | λ: 0.1 | Action: Integration test for database query extraction to XLSX and XML formats
import os
import sys
import json
import sqlite3
import subprocess
from pathlib import Path

def run_test():
    print("=== RUNNING EXPORT PRIMITIVES (XLSX & XML) INTEGRATION TEST ===")
    
    test_dir = Path(__file__).parent
    test_dir.mkdir(exist_ok=True)
    
    db_path = test_dir / "test_export_db.sqlite"
    json_path = test_dir / "db_users.json"
    xlsx_path = test_dir / "exported_users.xlsx"
    xml_path = test_dir / "exported_users.xml"
    
    # Clean up previous test files
    for p in [db_path, json_path, xlsx_path, xml_path]:
        if p.exists():
            try:
                p.unlink()
            except Exception:
                pass
                
    # 1. Setup sqlite local database and insert mock data
    conn = sqlite3.connect(db_path)
    cursor = conn.cursor()
    cursor.execute("""
        CREATE TABLE users (
            id INTEGER PRIMARY KEY,
            name TEXT,
            score REAL,
            city TEXT
        )
    """)
    cursor.executemany("""
        INSERT INTO users (id, name, score, city) VALUES (?, ?, ?, ?)
    """, [
        (1, "Alice", 95.5, "Paris"),
        (2, "Bob", 82.0, "Lyon"),
        (3, "Charlie", 78.5, "Marseille")
    ])
    conn.commit()
    conn.close()
    print(f"Created local SQLite database and inserted records at '{db_path}'")
    
    # Locate rust_muscle binary
    exe_ext = ".exe" if os.name == "nt" else ""
    binary_path = Path(__file__).parent.parent / "rust_muscle" / "target" / "debug" / f"rust_muscle{exe_ext}"
    
    if not binary_path.exists():
        print(f"Error: Compiled binary not found at '{binary_path}'")
        sys.exit(1)
        
    # 2. Run db.query to extract SQLite table into a JSON file
    conn_str = f"sqlite://{db_path}"
    query = "SELECT id, name, score, city FROM users ORDER BY id ASC"
    
    cmd_query = [
        str(binary_path),
        "db.query",
        "--connection_string", conn_str,
        "--query", query,
        "--destination", str(json_path)
    ]
    
    print(f"Executing Query: {' '.join(cmd_query)}")
    res_query = subprocess.run(cmd_query, capture_output=True, text=True, encoding="utf-8")
    print(f"Exit code: {res_query.returncode}")
    if res_query.returncode != 0:
        print(f"Stderr: {res_query.stderr}")
        sys.exit(1)
        
    assert json_path.exists(), "Query output JSON file was not created"
    
    # 3. Run data.to_xlsx to convert JSON to Excel (.xlsx)
    cmd_xlsx = [
        str(binary_path),
        "data.to_xlsx",
        "--source", str(json_path),
        "--destination", str(xlsx_path),
        "--sheet-name", "Utilisateurs"
    ]
    
    print(f"Executing Excel Export: {' '.join(cmd_xlsx)}")
    res_xlsx = subprocess.run(cmd_xlsx, capture_output=True, text=True, encoding="utf-8")
    print(f"Exit code: {res_xlsx.returncode}")
    if res_xlsx.returncode != 0:
        print(f"Stderr: {res_xlsx.stderr}")
        sys.exit(1)
        
    assert xlsx_path.exists(), "Excel output file was not created"
    
    # 4. Run data.json_to_xml to convert JSON to XML
    cmd_xml = [
        str(binary_path),
        "data.json_to_xml",
        "--source", str(json_path),
        "--destination", str(xml_path),
        "--root-element", "Company",
        "--row-element", "Employee"
    ]
    
    print(f"Executing XML Export: {' '.join(cmd_xml)}")
    res_xml = subprocess.run(cmd_xml, capture_output=True, text=True, encoding="utf-8")
    print(f"Exit code: {res_xml.returncode}")
    if res_xml.returncode != 0:
        print(f"Stderr: {res_xml.stderr}")
        sys.exit(1)
        
    assert xml_path.exists(), "XML output file was not created"
    
    # 5. Verify XML contents
    import xml.etree.ElementTree as ET
    tree = ET.parse(xml_path)
    root = tree.getroot()
    
    assert root.tag == "Company", f"Expected root tag 'Company', got '{root.tag}'"
    employees = root.findall("Employee")
    assert len(employees) == 3, f"Expected 3 Employee elements, got {len(employees)}"
    
    assert employees[0].find("name").text == "Alice"
    assert employees[1].find("name").text == "Bob"
    assert employees[2].find("name").text == "Charlie"
    assert employees[0].find("city").text == "Paris"
    print("XML Structure verification: SUCCESS")
    
    # 6. Verify Excel contents
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path)
    assert "Utilisateurs" in wb.sheetnames, "Expected sheet 'Utilisateurs' not found"
    ws = wb["Utilisateurs"]
    
    # Check headers
    headers = [cell.value for cell in ws[1]]
    assert headers == ["city", "id", "name", "score"], f"Unexpected headers: {headers}"
    
    # Check rows
    row2 = [cell.value for cell in ws[2]]
    row3 = [cell.value for cell in ws[3]]
    row4 = [cell.value for cell in ws[4]]
    
    # Verify values and types
    assert row2 == ["Paris", 1, "Alice", 95.5]
    assert row3 == ["Lyon", 2, "Bob", 82.0]
    assert row4 == ["Marseille", 3, "Charlie", 78.5]
    print("Excel Structure and Data Type verification: SUCCESS")
    
    print("\n[VERIFICATION_GATE]")
    print("- Invariant 14 [Excel XLSX Export Primitive]: SUCCESS")
    print("- Invariant 15 [XML Document Generation Primitive]: SUCCESS")
    
    # Clean up files (keep xlsx and xml for user inspection)
    for p in [db_path, json_path]:
        if p.exists():
            try:
                p.unlink()
            except Exception:
                pass
                
    print("\nINTEGRATION TEST COMPLETED SUCCESSFULLY!")

if __name__ == "__main__":
    run_test()
